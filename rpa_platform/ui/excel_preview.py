import os
from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem,
    QPushButton, QLabel, QComboBox, QLineEdit, QHeaderView,
    QAbstractItemView, QFrame, QSizePolicy,
)
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor, QFont, QBrush

from .theme import TXT1, TXT3, DIVIDER, ACCENT


class ExcelPreviewDialog(QDialog):
    def __init__(self, file_path: str = "", parent=None):
        super().__init__(parent)
        self.setWindowTitle("Vista previa — Excel")
        self.resize(900, 560)
        self._file_path = file_path
        self._wb = None
        self._setup_ui()

        if file_path and os.path.exists(file_path):
            QTimer.singleShot(100, self._load_file)

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        # Toolbar
        toolbar = QHBoxLayout()
        toolbar.setSpacing(8)

        self._lbl_file = QLabel("Sin archivo")
        self._lbl_file.setStyleSheet(f"color: {TXT1}; font-size: 12px; font-weight: 600;")
        self._lbl_file.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        toolbar.addWidget(self._lbl_file)

        toolbar.addWidget(QLabel("Hoja:"))
        self._sheet_combo = QComboBox()
        self._sheet_combo.setFixedWidth(150)
        self._sheet_combo.currentTextChanged.connect(self._on_sheet_changed)
        toolbar.addWidget(self._sheet_combo)

        toolbar.addWidget(QLabel("Buscar:"))
        self._search_box = QLineEdit()
        self._search_box.setFixedWidth(160)
        self._search_box.setPlaceholderText("Texto...")
        self._search_box.textChanged.connect(self._on_search)
        toolbar.addWidget(self._search_box)

        self._btn_close = QPushButton("Cerrar")
        self._btn_close.clicked.connect(self.close)
        toolbar.addWidget(self._btn_close)

        layout.addLayout(toolbar)

        sep = QFrame()
        sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet(f"background-color: {DIVIDER};")
        sep.setFixedHeight(1)
        layout.addWidget(sep)

        # Table
        self._table = QTableWidget()
        self._table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self._table.setAlternatingRowColors(True)
        self._table.setSelectionMode(QAbstractItemView.SelectionMode.ContiguousSelection)
        self._table.horizontalHeader().setDefaultSectionSize(100)
        self._table.verticalHeader().setDefaultSectionSize(22)
        self._table.setAlternatingRowColors(True)
        layout.addWidget(self._table)

        # Status bar
        self._status = QLabel("—")
        self._status.setStyleSheet(f"color: {TXT3}; font-size: 11px;")
        layout.addWidget(self._status)

    def load(self, file_path: str):
        self._file_path = file_path
        self._load_file()

    def _load_file(self):
        if not self._file_path or not os.path.exists(self._file_path):
            return
        try:
            import openpyxl
            self._wb = openpyxl.load_workbook(self._file_path, read_only=True, data_only=True)
            self._lbl_file.setText(os.path.basename(self._file_path))

            self._sheet_combo.blockSignals(True)
            self._sheet_combo.clear()
            self._sheet_combo.addItems(self._wb.sheetnames)
            self._sheet_combo.blockSignals(False)

            self._populate_table(self._wb.active)
        except Exception as e:
            self._status.setText(f"Error: {e}")

    def _on_sheet_changed(self, name: str):
        if self._wb and name in self._wb.sheetnames:
            self._populate_table(self._wb[name])

    def _populate_table(self, ws):
        self._table.clearContents()
        self._table.setRowCount(0)
        self._table.setColumnCount(0)

        # Read up to 500 rows, 50 cols for preview
        max_rows, max_cols = 500, 50
        rows = list(ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True))
        if not rows:
            self._status.setText("Hoja vacía")
            return

        n_cols = max(len(r) for r in rows)
        self._table.setRowCount(len(rows))
        self._table.setColumnCount(n_cols)

        # Column headers: A, B, C…
        from openpyxl.utils import get_column_letter
        self._table.setHorizontalHeaderLabels(
            [get_column_letter(i + 1) for i in range(n_cols)]
        )

        # Row headers: 1, 2, 3…
        self._table.setVerticalHeaderLabels([str(i + 1) for i in range(len(rows))])

        header_font = QFont()
        header_font.setBold(True)

        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                text = "" if val is None else str(val)
                item = QTableWidgetItem(text)
                item.setTextAlignment(Qt.AlignmentFlag.AlignVCenter | Qt.AlignmentFlag.AlignLeft)
                if r_idx == 0:
                    item.setFont(header_font)
                    item.setForeground(QBrush(QColor(ACCENT)))
                self._table.setItem(r_idx, c_idx, item)

        self._table.resizeColumnsToContents()
        # Cap column width
        for c in range(self._table.columnCount()):
            if self._table.columnWidth(c) > 200:
                self._table.setColumnWidth(c, 200)

        truncated = len(rows) >= max_rows
        self._status.setText(
            f"{len(rows)} filas × {n_cols} columnas"
            + (" (mostrando primeras 500)" if truncated else "")
        )

    def _on_search(self, text: str):
        text = text.strip().lower()
        for r in range(self._table.rowCount()):
            visible = False
            for c in range(self._table.columnCount()):
                item = self._table.item(r, c)
                if item and text in item.text().lower():
                    visible = True
                    break
            self._table.setRowHidden(r, not visible if text else False)

        if text:
            self._status.setText(
                f"Filtrando por «{text}» — "
                + str(sum(1 for r in range(self._table.rowCount())
                          if not self._table.isRowHidden(r)))
                + " filas visibles"
            )


def open_excel_preview(file_path: str, parent=None):
    dlg = ExcelPreviewDialog(file_path=file_path, parent=parent)
    dlg.exec()
