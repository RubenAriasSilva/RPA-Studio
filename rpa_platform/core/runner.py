import threading
import time
from typing import Dict, Any, Callable, Optional
from enum import Enum

from .workflow import Workflow, WorkflowStep
from actions.base import ActionResult, ActionStatus

# Loop-control action types the runner handles natively
LOOP_START_TYPES = ("loop_start", "loop", "excel_foreach_row")
LOOP_END_TYPE = "loop_end"


class RunnerState(Enum):
    IDLE = "idle"
    RUNNING = "running"
    PAUSED = "paused"
    STOPPED = "stopped"
    FINISHED = "finished"
    FAILED = "failed"      # a step asked to end the whole process with an error


class WorkflowRunner:
    def __init__(self):
        self.state = RunnerState.IDLE
        self._thread: Optional[threading.Thread] = None
        self._stop_event = threading.Event()
        self._pause_event = threading.Event()
        self._pause_event.set()
        # Keys pressed by "Mantener Tecla - Inicio" and not yet released, innermost last
        self._held_keys: list = []

        self.on_step_start: Optional[Callable] = None
        self.on_step_complete: Optional[Callable] = None
        self.on_log: Optional[Callable] = None
        self.on_finished: Optional[Callable] = None
        self.on_error: Optional[Callable] = None

    # ------------------------------------------------------------------ public

    def run(self, workflow: Workflow):
        if self.state == RunnerState.RUNNING:
            return
        self._stop_event.clear()
        self._pause_event.set()
        self.state = RunnerState.RUNNING
        self._thread = threading.Thread(target=self._execute, args=(workflow,), daemon=True)
        self._thread.start()

    def stop(self):
        self._stop_event.set()
        self._pause_event.set()   # unblock any paused wait so thread can exit

    def pause(self):
        if self.state == RunnerState.RUNNING:
            self._pause_event.clear()
            self.state = RunnerState.PAUSED

    def resume(self):
        if self.state == RunnerState.PAUSED:
            self._pause_event.set()
            self.state = RunnerState.RUNNING

    @property
    def is_running(self):
        return self.state in (RunnerState.RUNNING, RunnerState.PAUSED)

    # ----------------------------------------------------------------- private

    def _log(self, message: str, level: str = "info"):
        if self.on_log:
            self.on_log(message, level)

    def _release_keys(self):
        """Never leave a key stuck down in Windows once the workflow ends."""
        if not self._held_keys:
            return
        import pyautogui
        for keys in reversed(self._held_keys):
            for key in reversed(keys):
                try:
                    pyautogui.keyUp(key)
                except Exception:
                    pass
        self._log(f"  ⌨ Teclas soltadas al terminar: {', '.join('+'.join(k) for k in self._held_keys)}", "warning")
        self._held_keys.clear()

    def _execute(self, workflow: Workflow):
        self._held_keys.clear()
        try:
            self._execute_steps(workflow)
        finally:
            self._release_keys()

    def _execute_steps(self, workflow: Workflow):
        from actions import ACTION_REGISTRY

        variables: Dict[str, Any] = dict(workflow.variables)
        steps = workflow.steps
        total = len(steps)

        # loop_state maps loop_start step_index →
        #   for loop_start:      {"count": int, "max": int}
        #   for excel_foreach:   {"rows": list, "row_idx": int}
        loop_state: Dict[int, Dict] = {}

        self._log(f"▶  {workflow.name}  ({sum(1 for s in steps if workflow.is_step_active(s))} pasos activos)", "info")

        idx = 0
        while idx < total:

            # ── check stop ──────────────────────────────────────────────────
            if self._stop_event.is_set():
                self.state = RunnerState.STOPPED
                self._log("Detenido por el usuario", "warning")
                self._release_keys()
                if self.on_finished:
                    self.on_finished(variables)
                return

            # ── check pause (blocks here until resumed) ─────────────────────
            self._pause_event.wait()
            if self._stop_event.is_set():
                continue   # re-check stop after resume

            step = steps[idx]

            if not workflow.is_step_active(step):
                idx += 1
                continue

            action_cls = ACTION_REGISTRY.get(step.action_type)
            if not action_cls:
                self._log(f"[{idx+1}] Acción desconocida: {step.action_type}", "error")
                idx += 1
                continue

            # ── loop_start ───────────────────────────────────────────────────
            if step.action_type in ("loop_start", "loop"):
                if idx not in loop_state:
                    max_iter = int(step.params.get("iterations", 3))
                    loop_state[idx] = {"count": 0, "max": max_iter}
                state = loop_state[idx]
                if state["max"] <= 0:
                    self._log(f"[{idx+1}] 🔁 Loop — 0 iteraciones, saltando", "info")
                    if self.on_step_start:
                        self.on_step_start(idx, step)
                    if self.on_step_complete:
                        self.on_step_complete(idx, step, ActionResult.success(output="0 iteraciones"))
                    del loop_state[idx]
                    idx = self._find_loop_end(steps, idx) + 1
                    continue
                counter_var = step.params.get("counter_variable", "loop_index")
                variables[counter_var] = str(state["count"])
                self._log(f"[{idx+1}] 🔁 Loop — iteración {state['count'] + 1}/{state['max']}", "info")
                if self.on_step_start:
                    self.on_step_start(idx, step)
                if self.on_step_complete:
                    self.on_step_complete(idx, step, ActionResult.success(output=f"Iteración {state['count']+1}"))
                idx += 1
                continue

            # ── excel_foreach_row ────────────────────────────────────────────
            if step.action_type == "excel_foreach_row":
                if idx not in loop_state:
                    rows, err = self._load_excel_rows(step.params)
                    if err:
                        self._log(f"[{idx+1}] ✗ {err}", "error")
                        # skip to matching loop_end
                        idx = self._find_loop_end(steps, idx) + 1
                        continue
                    loop_state[idx] = {"rows": rows, "row_idx": 0}
                state = loop_state[idx]
                if state["row_idx"] >= len(state["rows"]):
                    self._log(f"[{idx+1}] 🔁 Excel loop terminado ({len(state['rows'])} filas)", "info")
                    del loop_state[idx]
                    idx = self._find_loop_end(steps, idx) + 1
                    continue
                self._set_excel_row_vars(variables, state["rows"][state["row_idx"]], step.params)
                self._log(
                    f"[{idx+1}] 🔁 Excel fila {state['row_idx'] + 1}/{len(state['rows'])}",
                    "info",
                )
                if self.on_step_start:
                    self.on_step_start(idx, step)
                if self.on_step_complete:
                    self.on_step_complete(idx, step, ActionResult.success(
                        output=f"Fila {state['row_idx'] + 1}/{len(state['rows'])}"
                    ))
                idx += 1
                continue

            # ── loop_end ─────────────────────────────────────────────────────
            if step.action_type == "loop_end":
                if self.on_step_start:
                    self.on_step_start(idx, step)

                start_idx = self._find_loop_start(steps, idx, loop_state)
                if start_idx is None:
                    self._log(f"[{idx+1}] ⚠ Fin de loop sin inicio", "warning")
                    if self.on_step_complete:
                        self.on_step_complete(idx, step, ActionResult.success(output="Fin de loop"))
                    idx += 1
                    continue

                start_step = steps[start_idx]
                state = loop_state[start_idx]

                if start_step.action_type in ("loop_start", "loop"):
                    state["count"] += 1
                    if state["count"] < state["max"]:
                        counter_var = start_step.params.get("counter_variable", "loop_index")
                        variables[counter_var] = str(state["count"])
                        if self.on_step_complete:
                            self.on_step_complete(idx, step, ActionResult.success(output=f"Iteración {state['count']}/{state['max']}"))
                        idx = start_idx + 1
                        continue
                    else:
                        self._log(f"  ✓ Loop completado ({state['max']} iteraciones)", "success")
                        del loop_state[start_idx]

                elif start_step.action_type == "excel_foreach_row":
                    state["row_idx"] += 1
                    if state["row_idx"] < len(state["rows"]):
                        self._set_excel_row_vars(variables, state["rows"][state["row_idx"]], start_step.params)
                        if self.on_step_complete:
                            self.on_step_complete(idx, step, ActionResult.success(
                                output=f"Fila {state['row_idx'] + 1}/{len(state['rows'])}"
                            ))
                        idx = start_idx + 1
                        continue
                    else:
                        self._log(f"  ✓ Excel loop completado ({len(state['rows'])} filas)", "success")
                        del loop_state[start_idx]

                if self.on_step_complete:
                    self.on_step_complete(idx, step, ActionResult.success(output="Loop terminado"))
                idx += 1
                continue

            # ── normal action ────────────────────────────────────────────────
            action = action_cls(params=step.params, step_id=step.step_id)
            context = {
                "variables": variables,
                "step_index": idx,
                "total_steps": total,
                "stop_event": self._stop_event,
                "pause_event": self._pause_event,
                "log": self._log,
                "held_keys": self._held_keys,
            }

            if self.on_step_start:
                self.on_step_start(idx, step)

            self._log(f"[{idx+1}/{total}] {action.get_summary()}", "info")

            try:
                result: ActionResult = action.execute(context)
            except Exception as e:
                result = ActionResult.failure(str(e))

            if result.variables:
                variables.update(result.variables)

            if result.status == ActionStatus.SUCCESS:
                self._log(f"  ✓ {result.output}", "success")
            else:
                self._log(f"  ✗ {result.error}", "error")
                if self.on_error:
                    self.on_error(idx, step, result.error)

            if self.on_step_complete:
                self.on_step_complete(idx, step, result)

            if result.abort:
                self.state = RunnerState.FAILED
                self._log(f"⛔  Proceso finalizado por error en el paso {idx+1}: {workflow.name}", "error")
                self._release_keys()
                if self.on_finished:
                    self.on_finished(variables)
                return

            idx += 1

        self.state = RunnerState.FINISHED
        self._log(f"✓  Workflow completado: {workflow.name}", "success")
        self._release_keys()
        if self.on_finished:
            self.on_finished(variables)

    # ----------------------------------------------------------------- helpers

    def _find_loop_end(self, steps, start_idx: int) -> int:
        depth = 0
        for i in range(start_idx, len(steps)):
            if steps[i].action_type in LOOP_START_TYPES:
                depth += 1
            elif steps[i].action_type == LOOP_END_TYPE:
                depth -= 1
                if depth == 0:
                    return i
        return len(steps) - 1

    def _find_loop_start(self, steps, end_idx: int, loop_state: Dict) -> Optional[int]:
        depth = 0
        for i in range(end_idx, -1, -1):
            if steps[i].action_type == LOOP_END_TYPE:
                depth += 1
            elif steps[i].action_type in LOOP_START_TYPES:
                depth -= 1
                if depth == 0 and i in loop_state:
                    return i
        return None

    def _load_excel_rows(self, params: Dict):
        import os
        try:
            import openpyxl
            path = params.get("file_path", "")
            if not path or not os.path.exists(path):
                return [], f"Archivo no encontrado: {path}"
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_name = params.get("sheet", "")
            ws = wb.active if not sheet_name else wb[sheet_name]
            start_row = max(1, int(params.get("start_row", 2)))
            end_row = int(params.get("end_row", 0))
            rows = []
            for row in ws.iter_rows(min_row=start_row, values_only=True):
                if end_row and len(rows) >= (end_row - start_row + 1):
                    break
                if all(v is None for v in row):
                    break
                rows.append([str(v) if v is not None else "" for v in row])
            wb.close()
            return rows, None
        except Exception as e:
            return [], str(e)

    def _set_excel_row_vars(self, variables: Dict, row: list, params: Dict):
        from openpyxl.utils import get_column_letter
        prefix = params.get("variable_prefix", "col_")
        col_map_str = params.get("column_map", "")

        # Always set col_A, col_B, col_C...
        for i, val in enumerate(row):
            letter = get_column_letter(i + 1)
            variables[f"{prefix}{letter}"] = val

        # Also apply named mapping if provided: "nombre=A,fecha=B"
        if col_map_str:
            for mapping in col_map_str.split(","):
                mapping = mapping.strip()
                if "=" in mapping:
                    var_name, col_letter = mapping.split("=", 1)
                    col_idx = 0
                    from openpyxl.utils import column_index_from_string
                    try:
                        col_idx = column_index_from_string(col_letter.strip().upper()) - 1
                    except Exception:
                        pass
                    if 0 <= col_idx < len(row):
                        variables[var_name.strip()] = row[col_idx]
