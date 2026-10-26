import os
import json
import re
from typing import Dict, Any, Optional, List, Tuple
from .base import BaseAction, ActionResult
from .windows import ON_TIMEOUT_FAIL, ON_TIMEOUT_CHOICES


# ---------- helpers ----------------------------------------------------------

def _open_wb(path: str, read_only: bool = False):
    import openpyxl
    if not os.path.exists(path):
        raise FileNotFoundError(f"Archivo no encontrado: {path}")
    return openpyxl.load_workbook(path, read_only=read_only, data_only=True)


def _get_sheet(wb, sheet_name: str):
    if not sheet_name or sheet_name.strip() == "":
        return wb.active
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    # Try by index
    try:
        idx = int(sheet_name)
        return wb.worksheets[idx]
    except (ValueError, IndexError):
        raise ValueError(f"Hoja '{sheet_name}' no encontrada. Hojas disponibles: {wb.sheetnames}")


def _col_letter_to_index(col: str) -> int:
    """'A' → 1, 'B' → 2, 'AA' → 27"""
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(col.upper())


def _parse_cell_ref(ref: str) -> Tuple[int, int]:
    """'B5' → (row=5, col=2)"""
    m = re.match(r"^([A-Za-z]+)(\d+)$", ref.strip())
    if not m:
        raise ValueError(f"Referencia de celda inválida: '{ref}'. Usa formato A1, B5, etc.")
    col = _col_letter_to_index(m.group(1))
    row = int(m.group(2))
    return row, col


def _cell_value_str(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    return str(v)


# ---------- actions ----------------------------------------------------------

class ExcelReadCellAction(BaseAction):
    action_type = "excel_read_cell"
    label = "Excel - Leer Celda"
    icon = "📗"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "cell": {"type": "str", "label": "Celda (ej: B5, A1)", "required": True, "default": "A1"},
        "output_variable": {"type": "str", "label": "Variable para el valor", "default": "excel_value"},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        path = params.get("file_path", "")
        sheet_name = params.get("sheet", "")
        cell_ref = params.get("cell", "A1").strip().upper()
        output_var = params.get("output_variable", "excel_value")

        try:
            wb = _open_wb(path, read_only=True)
            ws = _get_sheet(wb, sheet_name)
            row, col = _parse_cell_ref(cell_ref)
            value = _cell_value_str(ws.cell(row=row, column=col))
            wb.close()
            return ActionResult.success(
                output=f"{cell_ref} = «{value}»",
                variables={output_var: value},
            )
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        cell = self.params.get("cell", "A1")
        var = self.params.get("output_variable", "excel_value")
        return f"Excel leer {cell} -> {var}"


class ExcelReadRangeAction(BaseAction):
    action_type = "excel_read_range"
    label = "Excel - Leer Rango"
    icon = "📊"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "range": {"type": "str", "label": "Rango (ej: A1:D10, vacío = datos completos)", "default": ""},
        "output_format": {"type": "choice", "label": "Formato de salida", "default": "json",
                          "choices": ["json", "csv", "first_row_only"]},
        "output_variable": {"type": "str", "label": "Variable para los datos", "default": "excel_data"},
        "skip_empty_rows": {"type": "bool", "label": "Omitir filas vacías", "default": True},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        path = params.get("file_path", "")
        sheet_name = params.get("sheet", "")
        range_ref = params.get("range", "").strip().upper()
        fmt = params.get("output_format", "json")
        output_var = params.get("output_variable", "excel_data")
        skip_empty = params.get("skip_empty_rows", True)

        try:
            wb = _open_wb(path, read_only=True)
            ws = _get_sheet(wb, sheet_name)

            if range_ref:
                rows_iter = ws[range_ref]
                if not isinstance(rows_iter, tuple):
                    rows_iter = ((rows_iter,),)
                elif rows_iter and not isinstance(rows_iter[0], tuple):
                    rows_iter = (rows_iter,)
            else:
                rows_iter = list(ws.iter_rows())

            rows: List[List[str]] = []
            for row in rows_iter:
                values = [_cell_value_str(c) for c in row]
                if skip_empty and all(v == "" for v in values):
                    continue
                rows.append(values)

            wb.close()

            if fmt == "json":
                result_str = json.dumps(rows, ensure_ascii=False)
            elif fmt == "csv":
                result_str = "\n".join(",".join(r) for r in rows)
            elif fmt == "first_row_only":
                result_str = ",".join(rows[0]) if rows else ""

            return ActionResult.success(
                output=f"{len(rows)} filas leidas -> {output_var}",
                variables={output_var: result_str},
            )
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        rng = self.params.get("range", "") or "completo"
        return f"Excel leer rango {rng}"


class ExcelWriteCellAction(BaseAction):
    action_type = "excel_write_cell"
    label = "Excel - Escribir Celda"
    icon = "✏️"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "cell": {"type": "str", "label": "Celda (ej: B5)", "required": True, "default": "A1"},
        "value": {"type": "str", "label": "Valor a escribir", "default": ""},
        "from_variable": {
                    "type": "str",
                    "label": "Valor desde variable",
                    "default": "",
                },
        "auto_save": {"type": "bool", "label": "Guardar automáticamente", "default": True},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        variables = context.get("variables", {})

        path = params.get("file_path", "")
        sheet_name = params.get("sheet", "")
        cell_ref = params.get("cell", "A1").strip().upper()
        value = params.get("value", "")
        from_var = params.get("from_variable", "").strip()
        auto_save = params.get("auto_save", True)

        # Si se indicó una variable, se busca su valor por nombre
        if from_var:
            if from_var not in variables:
                return ActionResult.failure(f"La variable '{from_var}' no existe")
            final_value = variables[from_var]
        else:
            final_value = value

        try:
            wb = _open_wb(path, read_only=False)
            ws = _get_sheet(wb, sheet_name)
            row, col = _parse_cell_ref(cell_ref)
            ws.cell(row=row, column=col).value = final_value
            if auto_save:
                wb.save(path)
            wb.close()
            return ActionResult.success(output=f"{cell_ref} ← «{final_value}»")
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        cell = self.params.get("cell", "A1")
        val = str(self.params.get("value", ""))[:20]
        from_var = self.params.get("from_variable", "")
        if from_var:
            return f"Excel escribir {cell} = «{from_var}»"
        return f"Excel escribir {cell} = «{val}»"


class ExcelAppendRowAction(BaseAction):
    action_type = "excel_append_row"
    label = "Excel - Agregar Fila"
    icon = "➕"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "values": {"type": "str", "label": "Valores separados por coma (ej: {nombre},{fecha},{valor})",
                   "required": True},
        "start_column": {"type": "str", "label": "Columna inicial (ej: A)", "default": "A"},
        "auto_save": {"type": "bool", "label": "Guardar automáticamente", "default": True},
        "output_variable": {"type": "str", "label": "Variable para número de fila escrita", "default": "excel_row"},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        path = params.get("file_path", "")
        sheet_name = params.get("sheet", "")
        values_str = params.get("values", "")
        start_col_letter = params.get("start_column", "A").strip().upper()
        auto_save = params.get("auto_save", True)
        output_var = params.get("output_variable", "excel_row")

        if not values_str:
            return ActionResult.failure("Los valores no pueden estar vacíos")

        try:
            values = [v.strip() for v in values_str.split(",")]
            wb = _open_wb(path, read_only=False)
            ws = _get_sheet(wb, sheet_name)

            # Find next empty row in the start column
            start_col_idx = _col_letter_to_index(start_col_letter)
            next_row = 1
            for row in ws.iter_rows(min_col=start_col_idx, max_col=start_col_idx):
                if row[0].value is not None:
                    next_row = row[0].row + 1

            for offset, val in enumerate(values):
                ws.cell(row=next_row, column=start_col_idx + offset).value = val

            if auto_save:
                wb.save(path)
            wb.close()

            return ActionResult.success(
                output=f"Fila {next_row} agregada: {values[:4]}{'...' if len(values)>4 else ''}",
                variables={output_var: str(next_row)},
            )
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        vals = str(self.params.get("values", ""))[:30]
        return f"Excel agregar fila: {vals}"


class ExcelFormatRangeAction(BaseAction):
    action_type = "excel_format_range"
    label = "Excel - Dar Formato"
    icon = "🎨"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "range": {"type": "str", "label": "Rango (ej: A1:D1, A1)", "required": True, "default": "A1:A1"},
        "bold": {"type": "bool", "label": "Negrita", "default": False},
        "italic": {"type": "bool", "label": "Cursiva", "default": False},
        "font_size": {"type": "int", "label": "Tamaño de fuente (0 = sin cambio)", "default": 0},
        "font_color": {"type": "str", "label": "Color de fuente hex (ej: FF0000 = rojo, vacío = sin cambio)", "default": ""},
        "bg_color": {"type": "str", "label": "Color de fondo hex (ej: FFFF00 = amarillo, vacío = sin cambio)", "default": ""},
        "align": {"type": "choice", "label": "Alineación horizontal", "default": "none",
                  "choices": ["none", "left", "center", "right"]},
        "border": {"type": "bool", "label": "Borde fino alrededor", "default": False},
        "auto_save": {"type": "bool", "label": "Guardar automáticamente", "default": True},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        path = params.get("file_path", "")
        sheet_name = params.get("sheet", "")
        range_ref = params.get("range", "A1").strip().upper()

        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            wb = _open_wb(path, read_only=False)
            ws = _get_sheet(wb, sheet_name)

            bold = params.get("bold", False)
            italic = params.get("italic", False)
            font_size = int(params.get("font_size", 0))
            font_color = params.get("font_color", "").strip().lstrip("#")
            bg_color = params.get("bg_color", "").strip().lstrip("#")
            align = params.get("align", "none")
            border = params.get("border", False)

            cells = ws[range_ref]
            if not isinstance(cells, tuple):
                cells = ((cells,),)
            elif cells and not isinstance(cells[0], tuple):
                cells = (cells,)

            thin = Side(style="thin") if border else None
            border_obj = Border(left=thin, right=thin, top=thin, bottom=thin) if border else None

            for row in cells:
                for cell in row:
                    # Font
                    font_kwargs = {}
                    if bold:
                        font_kwargs["bold"] = True
                    if italic:
                        font_kwargs["italic"] = True
                    if font_size > 0:
                        font_kwargs["size"] = font_size
                    if font_color:
                        font_kwargs["color"] = font_color.upper()
                    if font_kwargs:
                        cell.font = Font(**font_kwargs)

                    # Fill
                    if bg_color:
                        cell.fill = PatternFill(
                            fill_type="solid",
                            fgColor=bg_color.upper(),
                        )

                    # Alignment
                    if align != "none":
                        cell.alignment = Alignment(horizontal=align)

                    # Border
                    if border_obj:
                        cell.border = border_obj

            if params.get("auto_save", True):
                wb.save(path)
            wb.close()

            return ActionResult.success(output=f"Formato aplicado a {range_ref}")
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        rng = self.params.get("range", "A1")
        parts = []
        if self.params.get("bold"):
            parts.append("negrita")
        if self.params.get("bg_color"):
            parts.append(f"fondo #{self.params['bg_color']}")
        return f"Excel formato {rng}" + (f": {', '.join(parts)}" if parts else "")


class ExcelFindValueAction(BaseAction):
    action_type = "excel_find_value"
    label = "Excel - Buscar Valor"
    icon = "🔍"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "search_value": {"type": "str", "label": "Valor a buscar", "required": True},
        "search_column": {"type": "str", "label": "Buscar solo en columna (ej: A, vacío = toda la hoja)", "default": ""},
        "exact_match": {"type": "bool", "label": "Coincidencia exacta", "default": False},
        "output_cell": {"type": "str", "label": "Variable para dirección de celda (ej: B5)", "default": "found_cell"},
        "output_row": {"type": "str", "label": "Variable para número de fila", "default": "found_row"},
        "output_value": {"type": "str", "label": "Variable para valor de la celda", "default": "found_value"},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        path = params.get("file_path", "")
        sheet_name = params.get("sheet", "")
        search = str(params.get("search_value", ""))
        search_col = params.get("search_column", "").strip().upper()
        exact = params.get("exact_match", False)
        out_cell = params.get("output_cell", "found_cell")
        out_row = params.get("output_row", "found_row")
        out_val = params.get("output_value", "found_value")

        try:
            wb = _open_wb(path, read_only=True)
            ws = _get_sheet(wb, sheet_name)

            from openpyxl.utils import get_column_letter

            col_filter = _col_letter_to_index(search_col) if search_col else None

            for row in ws.iter_rows():
                for cell in row:
                    if col_filter and cell.column != col_filter:
                        continue
                    val = _cell_value_str(cell)
                    match = (val == search) if exact else (search.lower() in val.lower())
                    if match:
                        col_letter = get_column_letter(cell.column)
                        cell_addr = f"{col_letter}{cell.row}"
                        wb.close()
                        return ActionResult.success(
                            output=f"Encontrado en {cell_addr}: «{val}»",
                            variables={
                                out_cell: cell_addr,
                                out_row: str(cell.row),
                                out_val: val,
                            },
                        )

            wb.close()
            return ActionResult.failure(f"Valor «{search}» no encontrado en la hoja")
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        val = str(self.params.get("search_value", ""))[:25]
        return f"Excel buscar «{val}»"


class ExcelSaveAction(BaseAction):
    action_type = "excel_save"
    label = "Excel - Guardar"
    icon = "💾"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo a guardar", "required": True},
        "save_as": {"type": "str", "label": "Guardar como (vacío = sobreescribir)", "default": ""},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        params = self.resolve_params(context.get("variables", {}))
        path = params.get("file_path", "")
        save_as = params.get("save_as", "").strip()

        try:
            wb = _open_wb(path, read_only=False)
            dest = save_as if save_as else path
            os.makedirs(os.path.dirname(os.path.abspath(dest)), exist_ok=True)
            wb.save(dest)
            wb.close()
            return ActionResult.success(output=f"Guardado: {os.path.basename(dest)}")
        except Exception as e:
            return ActionResult.failure(str(e))

    def get_summary(self) -> str:
        f = os.path.basename(self.params.get("file_path", ""))
        return f"Excel guardar: {f}"


class ExcelForEachRowAction(BaseAction):
    """
    Marks the start of an Excel row-iteration loop.
    The runner reads the file and advances one row per loop iteration.
    Place a 'Fin de Loop' action at the end of the steps you want to repeat.

    Variables set per iteration:
      {prefix}A, {prefix}B, {prefix}C …  (one per column)
      Plus any named mappings via column_map.
    """
    action_type = "excel_foreach_row"
    label = "Excel - Para Cada Fila"
    icon = "🔂"

    PARAM_SCHEMA = {
        "file_path": {"type": "file", "label": "Archivo Excel (.xlsx)", "required": True},
        "sheet": {"type": "str", "label": "Hoja (nombre o índice, vacío = primera)", "default": ""},
        "start_row": {"type": "int", "label": "Primera fila de datos (2 = saltar encabezado)", "default": 2},
        "end_row": {"type": "int", "label": "Última fila (0 = hasta la última con datos)", "default": 0},
        "variable_prefix": {"type": "str", "label": "Prefijo de variables de columna", "default": "col_"},
        "column_map": {
            "type": "str",
            "label": "Mapeo nombre=Columna (ej: nombre=A,fecha=B,valor=C)",
            "default": "",
        },
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        return ActionResult.success(output="Excel foreach — manejado por el runner")

    def get_summary(self) -> str:
        f = os.path.basename(self.params.get("file_path", ""))
        start = self.params.get("start_row", 2)
        return f"🔂 Para cada fila de {f} (desde fila {start})"


# ---------- open workbooks (live Excel, not files on disk) -------------------

_EXCEL_EXTS = (".xlsx", ".xlsm", ".xlsb", ".xls", ".xltx", ".xltm", ".csv")
# Status text Excel appends to the window title after the workbook name
_TITLE_STATUS = re.compile(
    r"\s*(\[[^\]]*\]|-\s*(Solo lectura|Read-Only|Modo de compatibilidad|Compatibility Mode|"
    r"Guardado|Saved|Guardando\S*|Saving\S*))\s*$",
    re.IGNORECASE,
)


def _workbook_name_from_title(title: str) -> str:
    """'reporte.xlsx [Modo de compatibilidad] - Excel' → 'reporte.xlsx'"""
    name, n = re.subn(r"\s+-\s+(Microsoft\s+)?Excel\s*$", "", title.strip(), flags=re.IGNORECASE)
    if not n:
        return ""  # plain "Excel": the start screen while it loads, not a workbook
    while True:
        stripped = _TITLE_STATUS.sub("", name)
        if stripped == name:
            break
        name = stripped
    return re.sub(r":\d+$", "", name).strip()  # 'reporte.xlsx:2' is a second window of the same book


def open_excel_workbooks() -> Dict[str, Tuple[str, Optional[int]]]:
    """Workbooks open in any Excel instance: {name in lower case: (full path or "", hwnd or None)}.

    Window titles catch every workbook, including unsaved ones and books opened
    in a window Excel reused. The Running Object Table adds the full path of
    every saved workbook, across all Excel instances.
    """
    from .windows import list_windows
    books: Dict[str, Tuple[str, Optional[int]]] = {}
    for w in list_windows():
        if w.class_name == "XLMAIN" and w.process == "excel.exe":
            name = _workbook_name_from_title(w.title)
            if name:
                books[name.lower()] = ("", w.hwnd)
    try:
        import pythoncom
        ctx = pythoncom.CreateBindCtx(0)
        for moniker in pythoncom.GetRunningObjectTable().EnumRunning():
            try:
                path = moniker.GetDisplayName(ctx, None)
            except Exception:
                continue
            if path.lower().endswith(_EXCEL_EXTS):
                key = os.path.basename(path).lower()
                books[key] = (path, books.get(key, ("", None))[1])
    except Exception:
        pass
    return books


class ExcelWaitWorkbookAction(BaseAction):
    action_type = "excel_wait_workbook"
    label = "Excel - Esperar Libro Nuevo"
    icon = "📗"

    PARAM_SCHEMA = {
        "name_contains": {"type": "str", "label": "Nombre del archivo contiene (vacío = cualquier libro nuevo)",
                          "placeholder": "ej. export, .xlsx"},
        "timeout": {"type": "float", "label": "Tiempo máximo de espera (seg)", "default": 120.0},
        "on_timeout": {"type": "choice", "label": "Si no se abre", "default": ON_TIMEOUT_FAIL,
                       "choices": ON_TIMEOUT_CHOICES},
        "focus": {"type": "bool", "label": "Enfocar el libro cuando se abra", "default": True},
        "output_variable": {"type": "str", "label": "Variable para el nombre del libro", "default": "excel_libro"},
        "path_variable": {"type": "str", "label": "Variable para la ruta completa (vacía si no está guardado)",
                          "default": "excel_ruta"},
    }

    def execute(self, context: Dict[str, Any]) -> ActionResult:
        from .windows import wait_until, timeout_result, focus_window, list_windows
        import pythoncom

        params = self.resolve_params(context.get("variables", {}))
        needle = str(params.get("name_contains", "")).strip().lower()
        timeout = max(0.0, float(params.get("timeout", 120.0)))
        name_var = str(params.get("output_variable", "excel_libro")).strip()
        path_var = str(params.get("path_variable", "excel_ruta")).strip()

        pythoncom.CoInitialize()  # the runner thread needs COM for the Running Object Table
        try:
            known = set(open_excel_workbooks())

            def find():
                for key, (path, hwnd) in open_excel_workbooks().items():
                    if key not in known and (not needle or needle in key):
                        return key, path, hwnd
                return None

            found, elapsed, stopped = wait_until(context, find, timeout, interval=0.5)
        finally:
            pythoncom.CoUninitialize()

        if stopped:
            return ActionResult.success(output="Espera de Excel interrumpida al detener el workflow")

        if found is None:
            variables = {k: "" for k in (name_var, path_var) if k}
            what = f"con «{params.get('name_contains', '').strip()}» " if needle else ""
            return timeout_result(params, f"No se abrió ningún libro de Excel {what}en {timeout:g} s", variables)

        key, path, hwnd = found
        name = os.path.basename(path) if path else key
        if hwnd is None:
            # Book known only through the ROT: find its window by title
            hwnd = next((w.hwnd for w in list_windows() if w.class_name == "XLMAIN"
                         and _workbook_name_from_title(w.title).lower() == key), None)
        if params.get("focus", True) and hwnd:
            focus_window(hwnd)
        variables = {}
        if name_var:
            variables[name_var] = name
        if path_var:
            variables[path_var] = path
        where = f" ({path})" if path else ""
        return ActionResult.success(output=f"Se abrió el libro «{name}»{where} tras {elapsed:.1f} s",
                                    variables=variables)

    def get_summary(self) -> str:
        needle = str(self.params.get("name_contains", "")).strip()
        target = f"«{needle}»" if needle else "nuevo"
        return f"Esperar libro de Excel {target} (máx {self.params.get('timeout', 120.0)}s)"
